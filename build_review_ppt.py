import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    # Load template to inherit layout & Anurag University theme
    template_path = 'mini_project)ppt/Team_No_8_Mini_Project_Presentation.pptx'
    prs = Presentation(template_path)
    
    # Clean existing slides
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[0]
        
    blank_layout = prs.slide_layouts[0] if len(prs.slide_layouts) > 0 else prs.slide_layouts[6]
    title_layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
    
    # Theme colors
    PRIMARY_BLUE = RGBColor(11, 59, 96)       # #0B3B60
    ACCENT_RED   = RGBColor(200, 16, 46)     # #C8102E
    DARK_TEXT    = RGBColor(30, 41, 59)      # #1E293B
    MUTED_TEXT   = RGBColor(71, 85, 105)     # #475569
    CARD_BG      = RGBColor(241, 245, 249)   # #F1F5F9
    BORDER_COLOR = RGBColor(203, 213, 225)   # #CBD5E1

    def add_header(slide, title_text):
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(8.8), Inches(0.7))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.name = "Georgia"
        p.font.color.rgb = PRIMARY_BLUE
        p.alignment = PP_ALIGN.CENTER

    # ----------------------------------------------------
    # SLIDE 1: Title Slide
    # ----------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    
    # Title
    t_box = s1.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(9.2), Inches(1.2))
    tf = t_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Customer Churn Prediction & Business Intelligence System"
    p.font.size = Pt(25)
    p.font.bold = True
    p.font.name = "Georgia"
    p.font.color.rgb = PRIMARY_BLUE
    p.alignment = PP_ALIGN.CENTER
    
    # Domain
    d_box = s1.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(9.2), Inches(0.6))
    tf_d = d_box.text_frame
    p_d = tf_d.paragraphs[0]
    p_d.text = "Domain: Business – Data Science & Telecom Analytics"
    p_d.font.size = Pt(17)
    p_d.font.bold = True
    p_d.font.color.rgb = ACCENT_RED
    p_d.alignment = PP_ALIGN.CENTER
    
    # Team Members (Left)
    team_box = s1.shapes.add_textbox(Inches(0.8), Inches(3.1), Inches(4.5), Inches(2.0))
    tf_t = team_box.text_frame
    tf_t.word_wrap = True
    p = tf_t.paragraphs[0]
    p.text = "Team No: 8"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    
    members = [
        "Bekkem Kalyan (23EG110C09)",
        "Eramolla Aryan Preetham (23EG110C21)",
        "Tirumala Peddinty Vijaya Krishna Sreekar (23EG110C60)"
    ]
    for m in members:
        p = tf_t.add_paragraph()
        p.text = m
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_TEXT
        
    # Guide Details (Right)
    guide_box = s1.shapes.add_textbox(Inches(5.6), Inches(3.1), Inches(4.4), Inches(2.0))
    tf_g = guide_box.text_frame
    tf_g.word_wrap = True
    p = tf_g.paragraphs[0]
    p.text = "Under the Guidance of"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    
    guide_lines = [
        "Mr. Biksham",
        "Assistant Professor",
        "Department of Data Science",
        "Anurag University"
    ]
    for gl in guide_lines:
        p = tf_g.add_paragraph()
        p.text = gl
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_TEXT

    # ----------------------------------------------------
    # SLIDE 2: Abstract
    # ----------------------------------------------------
    s2 = prs.slides.add_slide(title_layout)
    add_header(s2, "Abstract")
    
    abs_box = s2.shapes.add_textbox(Inches(0.9), Inches(1.3), Inches(8.8), Inches(3.8))
    tf = abs_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Customer churn is a critical business challenge in the telecom industry where retaining existing customers is 5 to 7 times more cost-effective than acquiring new ones. This project presents an end-to-end Machine Learning and Business Intelligence System designed to predict customer churn accurately and deliver actionable insights for customer retention."
    p.font.size = Pt(14)
    p.font.color.rgb = DARK_TEXT
    p.space_after = Pt(14)
    
    p2 = tf.add_paragraph()
    p2.text = "Key System Capabilities:"
    p2.font.size = Pt(14)
    p2.font.bold = True
    p2.font.color.rgb = PRIMARY_BLUE
    p2.space_after = Pt(6)
    
    abstract_bullets = [
        "Analyzes 7,043 customer records across 21 demographic, service, and billing attributes.",
        "Implements advanced preprocessing, outlier treatment, and SMOTE balancing to solve class imbalance.",
        "Trains and evaluates multiple algorithms (Random Forest, Logistic Regression, XGBoost) achieving 83.62% ROC-AUC.",
        "Integrates a SQLite database and an interactive Streamlit Web Dashboard for live predictions and SQL cohort analytics."
    ]
    for b in abstract_bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(4)

    # ----------------------------------------------------
    # SLIDE 3: Introduction
    # ----------------------------------------------------
    s3 = prs.slides.add_slide(title_layout)
    add_header(s3, "Introduction")
    
    intro_box = s3.shapes.add_textbox(Inches(0.9), Inches(1.2), Inches(8.8), Inches(4.0))
    tf = intro_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Problem Context & Industry Challenge"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_RED
    p.space_after = Pt(6)
    
    bullets = [
        "Telecom service providers experience heavy revenue loss due to customer contract cancellations (churn).",
        "Manual identification of at-risk customers is ineffective due to complex, multi-dimensional customer behavioral patterns.",
        "Unaddressed churn results in high acquisition overheads, reduced customer lifetime value (CLV), and diminished market share."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(4)
        
    p = tf.add_paragraph()
    p.text = "Proposed Solution & Data Flow"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    p.space_before = Pt(10)
    p.space_after = Pt(6)
    
    flow_bullets = [
        "Data Ingestion → Clean & Preprocess → Feature Engineering (SMOTE) → Multi-Model ML Training → SQL Analytics → Streamlit UI.",
        "Enables telecom managers to proactively intervene with targeted discounts, tenure perks, and service upgrades before churn occurs."
    ]
    for b in flow_bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(4)

    # ----------------------------------------------------
    # SLIDE 4: Objectives (Max 2 to 4 bullet points as per rule!)
    # ----------------------------------------------------
    s4 = prs.slides.add_slide(title_layout)
    add_header(s4, "Objectives")
    
    obj_box = s4.shapes.add_textbox(Inches(0.9), Inches(1.3), Inches(8.8), Inches(3.8))
    tf = obj_box.text_frame
    tf.word_wrap = True
    
    objectives = [
        "1. End-to-End Data Pipeline & Feature Engineering: Ingest, clean, and preprocess the IBM Telco Churn dataset, engineering composite risk metrics and applying SMOTE to address class imbalance.",
        "2. Predictive Model Optimization: Train and validate Machine Learning classifiers (Random Forest, XGBoost, Logistic Regression) to accurately identify churn patterns with high ROC-AUC and Recall.",
        "3. Financial & Cohort Analytics Engine: Build a robust SQL-powered analytics database to calculate monthly revenue at risk and isolate high-churn customer segments.",
        "4. Interactive Decision Support Dashboard: Deploy a real-time Streamlit web dashboard for instant single/batch churn prediction and executive reporting."
    ]
    
    for i, obj in enumerate(objectives):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = obj
        p.font.size = Pt(13.5)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(12)

    # ----------------------------------------------------
    # SLIDE 5: Literature Survey (Table Format)
    # ----------------------------------------------------
    s5 = prs.slides.add_slide(title_layout)
    add_header(s5, "Literature Survey")
    
    rows, cols = 6, 6
    left = Inches(0.5)
    top = Inches(1.2)
    width = Inches(9.5)
    height = Inches(3.8)
    
    table_shape = s5.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    
    # Set column widths
    table.columns[0].width = Inches(0.6)  # S.No
    table.columns[1].width = Inches(1.8)  # Author & Year
    table.columns[2].width = Inches(2.0)  # Method Used
    table.columns[3].width = Inches(1.8)  # Proposed Model
    table.columns[4].width = Inches(1.9)  # Drawbacks
    table.columns[5].width = Inches(1.4)  # Source
    
    headers = ["S.No", "Author & Year", "Method Used", "Proposed Model", "Drawbacks", "Source"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(10.5)
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = PP_ALIGN.CENTER
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY_BLUE
        
    lit_data = [
        ("1", "Ning Lu et al. (2014)", "Boosting with Customer Segmentation", "Telecom churn prediction via Boosting", "High computational cost; lacks real-time UI", "IEEE Trans. Ind. Inf."),
        ("2", "A. K. Ahmad et al. (2019)", "ML + Social Network Analysis", "DT, Random Forest, XGBoost", "Requires heavy infrastructure & big data pipelines", "J. of Big Data"),
        ("3", "Irfan Ullah et al. (2019)", "Random Forest + K-Means", "RF with customer clustering", "Complex multi-stage setup; no SQL cohort engine", "IEEE Access"),
        ("4", "Shimaa Ouf et al. (2024)", "Hybrid ML Framework", "Ensemble Stacking Classifiers", "High inference latency; no interactive retention UI", "Springer"),
        ("5", "Rohit Sharma et al. (2025)", "EDA + Multi-Model ML", "RF, SVM, Logistic Regression", "Single dataset evaluation; no business risk quantification", "IEEE ICONAT")
    ]
    
    for row_idx, data in enumerate(lit_data, start=1):
        for col_idx, text in enumerate(data):
            cell = table.cell(row_idx, col_idx)
            cell.text = text
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(9.5)
                p.font.color.rgb = DARK_TEXT

    # ----------------------------------------------------
    # SLIDE 6: Existing Work and Limitations
    # ----------------------------------------------------
    s6 = prs.slides.add_slide(title_layout)
    add_header(s6, "Existing Work and Limitations")
    
    box = s6.shapes.add_textbox(Inches(0.9), Inches(1.2), Inches(8.8), Inches(4.0))
    tf = box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Current / Existing Industry Approaches"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    p.space_after = Pt(4)
    
    existing_pts = [
        "Traditional Rule-Based Analytics: Many telecom providers rely on static heuristic threshold rules (e.g., flag accounts after 2 missed bills).",
        "Isolated Model Experiments: Prior studies focus exclusively on model training accuracy without deploying end-to-end usable business applications."
    ]
    for pt in existing_pts:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12.5)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(3)
        
    p = tf.add_paragraph()
    p.text = "Key Limitations of Existing Systems"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_RED
    p.space_before = Pt(8)
    p.space_after = Pt(4)
    
    limitations = [
        "Neglect of Class Imbalance: Churn datasets typically have a ~73:27 majority-minority split, causing standard models to be biased toward non-churners.",
        "Lack of Revenue Quantification: Existing solutions predict binary labels (0/1) but fail to quantify the direct financial revenue at risk ($/month).",
        "Absence of Real-Time Decision Support: Most models operate in offline batch scripts without an accessible web interface for customer support teams."
    ]
    for lim in limitations:
        p = tf.add_paragraph()
        p.text = "• " + lim
        p.font.size = Pt(12.5)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(3)

    # ----------------------------------------------------
    # SLIDE 7: Proposed Work and Advantages
    # ----------------------------------------------------
    s7 = prs.slides.add_slide(title_layout)
    add_header(s7, "Proposed Work and Advantages")
    
    box = s7.shapes.add_textbox(Inches(0.9), Inches(1.2), Inches(8.8), Inches(4.0))
    tf = box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Our Proposed Methodology"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    p.space_after = Pt(4)
    
    proposed_pts = [
        "End-to-End Automated Pipeline: Comprehensive orchestration from raw CSV ingestion to automated EDA, SMOTE balancing, SQL storage, and web UI.",
        "Advanced Feature Engineering: Creation of domain-specific features including ChargesPerTenure, ServiceCount, and a composite RiskScore.",
        "Multi-Model Comparative Benchmarking: Rigorous evaluation of Random Forest, XGBoost, and Logistic Regression with cross-validation."
    ]
    for pt in proposed_pts:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12.5)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(3)
        
    p = tf.add_paragraph()
    p.text = "Key Advantages of the Proposed System"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = RGBColor(4, 120, 87)
    p.space_before = Pt(8)
    p.space_after = Pt(4)
    
    advantages = [
        "High Predictive Reliability: Achieved 83.62% ROC-AUC and 70.43% churn recall, effectively capturing at-risk subscribers.",
        "Quantified Financial Impact: SQL analytics calculates exact revenue at risk ($138,529/mo) to prioritize high-value customer retention.",
        "Operational Ready: Includes an intuitive Streamlit dashboard allowing non-technical managers to run instant real-time predictions."
    ]
    for adv in advantages:
        p = tf.add_paragraph()
        p.text = "• " + adv
        p.font.size = Pt(12.5)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(3)

    # ----------------------------------------------------
    # SLIDE 8: Software & Hardware Requirements
    # ----------------------------------------------------
    s8 = prs.slides.add_slide(title_layout)
    add_header(s8, "Software and Hardware Requirements")
    
    # Software Requirements (Left Column)
    s_box = s8.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(4.3), Inches(3.9))
    tf_s = s_box.text_frame
    tf_s.word_wrap = True
    p = tf_s.paragraphs[0]
    p.text = "Software Requirements"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    p.space_after = Pt(6)
    
    sw_reqs = [
        "Operating System: Windows 10 / 11 / Linux",
        "Programming Language: Python 3.9+",
        "ML & Data Libraries: Scikit-learn, XGBoost, Imbalanced-learn (SMOTE), Pandas, NumPy",
        "Database Engine: SQLite3 & SQLAlchemy",
        "Visualization & Web UI: Streamlit, Matplotlib, Seaborn, Plotly",
        "Documentation & Reports: python-pptx, python-docx, ReportLab"
    ]
    for req in sw_reqs:
        p = tf_s.add_paragraph()
        p.text = "• " + req
        p.font.size = Pt(12)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(4)
        
    # Hardware Requirements (Right Column)
    h_box = s8.shapes.add_textbox(Inches(5.4), Inches(1.2), Inches(4.3), Inches(3.9))
    tf_h = h_box.text_frame
    tf_h.word_wrap = True
    p = tf_h.paragraphs[0]
    p.text = "Hardware Requirements"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_RED
    p.space_after = Pt(6)
    
    hw_reqs = [
        "Processor: Intel Core i5 / AMD Ryzen 5 or higher",
        "RAM: Minimum 8 GB (16 GB Recommended)",
        "Storage: 500 MB free disk space for dataset, database & models",
        "Display: Standard Full HD (1920x1080) resolution",
        "Connectivity: Active Internet connection for initial package installation and data retrieval"
    ]
    for req in hw_reqs:
        p = tf_h.add_paragraph()
        p.text = "• " + req
        p.font.size = Pt(12)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(4)

    # ----------------------------------------------------
    # SLIDE 9: System Architecture
    # ----------------------------------------------------
    s9 = prs.slides.add_slide(title_layout)
    add_header(s9, "System Architecture")
    
    arch_img_path = 'reports/figures/system_architecture_diagram.png'
    if os.path.exists(arch_img_path):
        s9.shapes.add_picture(arch_img_path, Inches(0.8), Inches(1.15), Inches(8.9), Inches(4.1))

    # ----------------------------------------------------
    # SLIDE 10: Modules Breakdown
    # ----------------------------------------------------
    s10 = prs.slides.add_slide(title_layout)
    add_header(s10, "Modules")
    
    m_box = s10.shapes.add_textbox(Inches(0.9), Inches(1.2), Inches(8.8), Inches(4.0))
    tf = m_box.text_frame
    tf.word_wrap = True
    
    modules = [
        ("Module 1: Data Ingestion & Preprocessing", "Handles raw CSV ingestion (7,043 rows), cleans whitespace/missing values in TotalCharges, encodes categorical variables, and performs standard feature scaling."),
        ("Module 2: Feature Engineering & SMOTE Resampling", "Derives predictive domain features (ChargesPerTenure, ServiceCount, RiskScore) and balances minority churn cases using SMOTE oversampling."),
        ("Module 3: Model Training & Evaluation Engine", "Trains Logistic Regression, Random Forest, and XGBoost with hyperparameter tuning; exports evaluation metrics (AUC, F1, Recall, Confusion Matrices)."),
        ("Module 4: SQL Business Intelligence & Cohort Analytics", "Stores cleaned data into SQLite (churn.db) and executes 10 strategic SQL queries to evaluate cohort churn and calculate revenue at risk ($138K/mo)."),
        ("Module 5: Interactive Web Dashboard (Streamlit)", "Provides a 5-page interactive web UI for real-time customer risk prediction, live database exploration, interactive charts, and PDF/Word report exports.")
    ]
    
    for i, (m_title, m_desc) in enumerate(modules):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "• " + m_title + ": "
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_BLUE
        
        run = p.add_run()
        run.text = m_desc
        run.font.bold = False
        run.font.color.rgb = DARK_TEXT
        p.space_after = Pt(5)

    # ----------------------------------------------------
    # SLIDE 11: Design (System Design & Functional Workflow)
    # ----------------------------------------------------
    s11 = prs.slides.add_slide(title_layout)
    add_header(s11, "System Design & Workflow")
    
    box = s11.shapes.add_textbox(Inches(0.9), Inches(1.2), Inches(8.8), Inches(4.0))
    tf = box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Functional Design & Layered Architecture"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE
    p.space_after = Pt(4)
    
    design_pts = [
        "Presentation Layer: Streamlit Dashboard UI receiving real-time user inputs and rendering visual analytics.",
        "Application / Service Layer: Python inference pipeline loading serialized models (`best_model.pkl`) and running live probability scoring.",
        "Data & Persistence Layer: SQLite Relational Database (`churn.db`) managing relational customer records, service contracts, and historical churn outcomes."
    ]
    for pt in design_pts:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12.5)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(4)
        
    p = tf.add_paragraph()
    p.text = "Design Specifications & Quality Attributes"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_RED
    p.space_before = Pt(8)
    p.space_after = Pt(4)
    
    spec_pts = [
        "Modularity: Decoupled architecture allowing independent model updates or data pipeline extensions.",
        "Responsiveness & Low Latency: Single customer churn predictions generated in < 50ms.",
        "Reliability & Maintainability: Automated unit checks, structured logging, and persistent artifact generation."
    ]
    for pt in spec_pts:
        p = tf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12.5)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(4)

    # ----------------------------------------------------
    # SLIDE 12: Design - UML, ER & DFD Diagrams
    # ----------------------------------------------------
    s12 = prs.slides.add_slide(title_layout)
    add_header(s12, "Design - UML, ER & DFD Diagrams")
    
    # Place 3 diagram preview thumbnails
    uml_path = 'reports/figures/uml_diagram.png'
    er_path  = 'reports/figures/er_diagram.png'
    dfd_path = 'reports/figures/dfd_diagram.png'
    
    if os.path.exists(uml_path) and os.path.exists(er_path):
        s12.shapes.add_picture(uml_path, Inches(0.5), Inches(1.15), Inches(4.5), Inches(1.9))
        s12.shapes.add_picture(dfd_path, Inches(5.3), Inches(1.15), Inches(4.5), Inches(1.9))
        s12.shapes.add_picture(er_path, Inches(2.6), Inches(3.15), Inches(5.0), Inches(2.0))

    # ----------------------------------------------------
    # SLIDE 13: Thank You
    # ----------------------------------------------------
    s13 = prs.slides.add_slide(blank_layout)
    
    t_box = s13.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.8), Inches(1.2))
    tf = t_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Thank You!"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.name = "Georgia"
    p.font.color.rgb = PRIMARY_BLUE
    p.alignment = PP_ALIGN.CENTER
    
    sub_box = s13.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(8.8), Inches(1.5))
    tf_sub = sub_box.text_frame
    p = tf_sub.paragraphs[0]
    p.text = "Questions & Feedback"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_RED
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf_sub.add_paragraph()
    p2.text = "Team No: 8  |  Department of Data Science  |  Anurag University"
    p2.font.size = Pt(14)
    p2.font.color.rgb = DARK_TEXT
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(8)

    # Save output
    output_path = 'Customer_Churn_Review_Presentation.pptx'
    prs.save(output_path)
    # Also save to main pptx location
    prs.save('Team_No_8_Mini_Project_Presentation.pptx')
    prs.save('mini_project)ppt/Team_No_8_Mini_Project_Presentation.pptx')
    print(f"Presentation successfully created and saved to {output_path}")

if __name__ == '__main__':
    create_presentation()
