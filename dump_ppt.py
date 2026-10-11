import pptx

prs = pptx.Presentation('Team_No_8_Mini_Project_Presentation.pptx')
for idx, slide in enumerate(prs.slides):
    print(f"\n=================== SLIDE {idx+1} ===================")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if p.text.strip():
                    print(f"- {p.text.strip()}")
        elif shape.has_table:
            for row in shape.table.rows:
                print(" | ".join(cell.text.strip() for cell in row.cells))
        else:
            print(f"[SHAPE]: {shape.name} (type: {shape.shape_type})")
