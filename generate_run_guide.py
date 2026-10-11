import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Set background color of a cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding of a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table):
    """Apply horizontal slate-grey borders."""
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    
    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'single')
    top.set(qn('w:sz'), '4')
    top.set(qn('w:space'), '0')
    top.set(qn('w:color'), '94A3B8')
    borders.append(top)
    
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '0')
    bottom.set(qn('w:color'), '64748B')
    borders.append(bottom)
    
    insideH = OxmlElement('w:insideH')
    insideH.set(qn('w:val'), 'single')
    insideH.set(qn('w:sz'), '4')
    insideH.set(qn('w:space'), '0')
    insideH.set(qn('w:color'), 'CBD5E1')
    borders.append(insideH)
    
    for border_name in ['left', 'right', 'insideV']:
        node = OxmlElement(f'w:{border_name}')
        node.set(qn('w:val'), 'none')
        borders.append(node)
        
    tblPr.append(borders)

def format_run(run, font_name="Calibri", size_pt=11, bold=False, italic=False, color_rgb=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    format_run(run, font_name="Arial", size_pt=16, bold=True, color_rgb=RGBColor(15, 23, 42))
    
    # Bottom accent line
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), '7C3AED') # Violet accent
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    format_run(run, font_name="Arial", size_pt=13, bold=True, color_rgb=RGBColor(124, 58, 237))
    return p

def add_body_text(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        format_run(r_prefix, font_name="Calibri", size_pt=11, bold=True, color_rgb=RGBColor(15, 23, 42))
    run = p.add_run(text)
    format_run(run, font_name="Calibri", size_pt=11, bold=False, italic=italic, color_rgb=RGBColor(51, 65, 85))
    return p

def add_bullet_point(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        format_run(r_prefix, font_name="Calibri", size_pt=11, bold=True, color_rgb=RGBColor(15, 23, 42))
    run = p.add_run(text)
    format_run(run, font_name="Calibri", size_pt=11, color_rgb=RGBColor(51, 65, 85))
    return p

def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    
    # Wrap in table for border/background
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F1F5F9") # slate-100
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    # Add border
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ['top', 'bottom', 'left', 'right']:
        node = OxmlElement(f'w:{side}')
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), '4')
        node.set(qn('w:color'), 'E2E8F0')
        tcBorders.append(node)
    tcPr.append(tcBorders)
    
    p_cell = cell.paragraphs[0]
    p_cell.paragraph_format.space_after = Pt(0)
    p_cell.paragraph_format.line_spacing = 1.15
    for i, line in enumerate(lines):
        r_line = p_cell.add_run(line + ("\n" if i < len(lines) - 1 else ""))
        format_run(r_line, font_name="Consolas", size_pt=9.5, color_rgb=RGBColor(15, 23, 42))
        
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_callout_box(doc, text, title="IMPORTANT NOTE"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "FAF5FF") # purple-50
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24')
    left.set(qn('w:color'), '7C3AED')
    tcBorders.append(left)
    for side in ['top', 'bottom', 'right']:
        node = OxmlElement(f'w:{side}')
        node.set(qn('w:val'), 'none')
        tcBorders.append(node)
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"★ {title}\n")
    format_run(r_title, font_name="Arial", size_pt=10, bold=True, color_rgb=RGBColor(109, 40, 217))
    
    r_text = p.add_run(text)
    format_run(r_text, font_name="Calibri", size_pt=10, italic=True, color_rgb=RGBColor(71, 85, 105))
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def main():
    print("Generating Run Guide...")
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Title
    for _ in range(2):
        doc.add_paragraph()
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_t = p_title.add_run("Setup & Execution Guide\n")
    format_run(run_t, font_name="Arial", size_pt=24, bold=True, color_rgb=RGBColor(15, 23, 42))
    run_st = p_title.add_run("Customer Churn Prediction System & CustomerPulse AI")
    format_run(run_st, font_name="Calibri", size_pt=14, italic=True, color_rgb=RGBColor(109, 40, 217))
    
    for _ in range(4):
        doc.add_paragraph()
        
    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run("PREPARED FOR:\tStandard Users & Development Teams\n")
    format_run(r_meta, font_name="Calibri", size_pt=10, color_rgb=RGBColor(100, 116, 139))
    r_meta2 = p_meta.add_run("LOCATION:\t\tDesktop (C:\\Users\\tpvij\\Desktop)\n")
    format_run(r_meta2, font_name="Calibri", size_pt=10, color_rgb=RGBColor(100, 116, 139))
    r_meta3 = p_meta.add_run("DATE:\t\tJune 29, 2026\n")
    format_run(r_meta3, font_name="Calibri", size_pt=10, color_rgb=RGBColor(100, 116, 139))
    
    doc.add_page_break()
    
    # Section 1: Overview
    add_heading_1(doc, "1. Project Directory Structure")
    add_body_text(doc, "The system is structured as an end-to-end Machine Learning pipeline with SQL support and an interactive web dashboard. Below is an explanation of the core files and folders in this workspace:")
    
    add_bullet_point(doc, "The master orchestration pipeline. Running this script downloads/generates the datasets, processes features, writes to SQL, fits all models, and saves serialized artifacts to disk.", "main.py: ")
    add_bullet_point(doc, "Exposes model comparative stats, feature importances, and SQL summaries in the console.", "show_results.py: ")
    add_bullet_point(doc, "Contains the raw and processed CSV folders, as well as the SQLite database (churn.db).", "data/: ")
    add_bullet_point(doc, "Stores serialized pickle files of trained models (best_model.pkl, logistic_regression.pkl, xgboost.pkl, random_forest.pkl), scaler models, encoders, and performance scores.", "models/: ")
    add_bullet_point(doc, "Contains the Streamlit web dashboard application (app.py) providing live KPIs and forecasting tabs.", "dashboard/: ")
    add_bullet_point(doc, "Contains 10 custom SQL analytics queries used to compute metrics such as Revenue-at-Risk.", "sql/: ")
    add_bullet_point(doc, "Holds the core code modules (preprocessing, feature engineering, modeling, SQL ingestion, utilities).", "src/: ")
    add_bullet_point(doc, "Holds the generated PNG plots from the EDA and evaluation cycles.", "reports/: ")
    add_bullet_point(doc, "Contains requirements.txt for setting up python environments.", "requirements.txt: ")
    add_bullet_point(doc, "The newly created enterprise SaaS structure containing PostgreSQL DDL schemas, docker compose scripts, and microservice skeletons.", "customerpulse-ai/: ")

    # Section 2: Prerequisites
    add_heading_1(doc, "2. System Prerequisites & Setup")
    add_body_text(doc, "Before running the project, guarantee that Python 3.9+ is installed. Follow these step-by-step commands to set up the local virtual environment:")
    
    add_heading_2(doc, "2.1  Create Virtual Environment")
    add_body_text(doc, "Open PowerShell or Command Prompt, navigate to the project directory, and create a virtual environment to isolate project dependencies:")
    add_code_block(doc, [
        "cd C:\\Users\\tpvij\\Desktop\\Customer_Churn_Prediction_System",
        "python -m venv venv"
    ])
    
    add_heading_2(doc, "2.2  Activate Virtual Environment")
    add_body_text(doc, "Activate the environment depending on your operating system:")
    add_code_block(doc, [
        "# Windows (PowerShell)",
        "venv\\Scripts\\Activate.ps1",
        "",
        "# Windows (CMD)",
        "venv\\Scripts\\activate.bat",
        "",
        "# Linux / macOS",
        "source venv/bin/activate"
    ])
    
    # Section 3: Installation
    add_heading_1(doc, "3. Package Installation")
    add_body_text(doc, "With the virtual environment activated, run pip to install the required data science and web libraries:")
    add_code_block(doc, [
        "pip install -r requirements.txt"
    ])
    add_body_text(doc, "If you want to run the SaaS reports generator module, also install python-docx:")
    add_code_block(doc, [
        "pip install python-docx"
    ])

    # Section 4: Execution
    add_heading_1(doc, "4. Step-by-Step Project Execution")
    add_body_text(doc, "Executing the project involves two main tasks: running the ML pipelines and launching the dashboard.")
    
    add_heading_2(doc, "4.1  Run the Machine Learning Pipeline")
    add_body_text(doc, "Execute the master orchestrator script. This script performs data cleaning, caps outliers, creates engineered features, oversamples minority classes using SMOTE, populates the SQL database, trains standard classifiers, selects the best model, and saves all metrics and charts:")
    add_code_block(doc, [
        "python main.py"
    ])
    add_body_text(doc, "Verify the outputs inside models/ and reports/folders. Run the results summary helper to check stats in your console:")
    add_code_block(doc, [
        "python show_results.py"
    ])
    
    add_heading_2(doc, "4.2  Launch the Streamlit Dashboard")
    add_body_text(doc, "Start the interactive web dashboard to explore the KPIs, forecast client risk, and view EDA insights:")
    add_code_block(doc, [
        "streamlit run dashboard/app.py"
    ])
    add_body_text(doc, "Streamlit will compile and launch the dashboard in your default browser at http://localhost:8501.")
    
    add_heading_2(doc, "4.3  Run local Enterprise SaaS Containers")
    add_body_text(doc, "For developers who want to inspect the newly scaffolded CustomerPulse AI multi-tenant services, boot up the local Docker container orchestration stack (ensure Docker Desktop is running):")
    add_code_block(doc, [
        "cd customerpulse-ai",
        "docker-compose up --build"
    ])

    # Section 5: Troubleshooting
    add_heading_1(doc, "5. Troubleshooting & Pitfalls")
    add_bullet_point(doc, "If python or pip commands are not found, check that Python is added to your Windows PATH variable.", "CommandNotFound: ")
    add_bullet_point(doc, "If executing show_results.py or app.py throws missing pickle errors, run 'python main.py' first to serialize and cache the models to disk.", "Missing Model Files: ")
    add_bullet_point(doc, "Ensure that your execution policy allows scripts to run: 'Set-ExecutionPolicy -ExecutionPolicy RemoteSigned -Scope Process'.", "PowerShell Execution Policy: ")
    add_bullet_point(doc, "If port 8501 is occupied, Streamlit will request another port automatically, or you can specify one: 'streamlit run dashboard/app.py --server.port 8080'.", "Port Conflict: ")
    
    add_callout_box(doc,
                    "Always activate the 'venv' virtual environment before executing scripts or loading dashboard interfaces. This keeps dependencies from leaking or conflicting with global site-packages.",
                    "BEST PRACTICE REQUIREMENT")

    # Save
    desktop_path = "C:\\Users\\tpvij\\Desktop\\Customer_Churn_Prediction_System_Run_Guide.docx"
    print(f"Saving Run Guide to: {desktop_path}")
    doc.save(desktop_path)
    print("Run Guide saved successfully!")

if __name__ == "__main__":
    main()
