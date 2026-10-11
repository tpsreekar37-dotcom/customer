import pptx

prs = pptx.Presentation('Team_No_8_Mini_Project_Presentation.pptx')
print(f"Total slides: {len(prs.slides)}")
for idx, slide in enumerate(prs.slides):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text:
            texts.append(shape.text_frame.text.strip())
    first_text = texts[0] if texts else "EMPTY"
    first_line = first_text.split('\n')[0]
    print(f"Slide {idx+1}: {first_line}")
