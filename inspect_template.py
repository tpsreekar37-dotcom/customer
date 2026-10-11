import pptx

prs = pptx.Presentation('mini_project)ppt/Team_No_8_Mini_Project_Presentation.pptx')
print("Slide width:", prs.slide_width, "height:", prs.slide_height)
for idx, slide in enumerate(prs.slides):
    print(f"\n--- Slide {idx+1} (layout: {slide.slide_layout.name}) ---")
    for s in slide.shapes:
        print(f"  Shape: {s.name} | Type: {s.shape_type} | Pos: ({s.left}, {s.top}, {s.width}, {s.height})")
