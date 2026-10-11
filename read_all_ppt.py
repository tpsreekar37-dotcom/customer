import pptx
import sys

# Set stdout encoding
sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation('mini_project)ppt/Team_No_8_Mini_Project_Presentation.pptx')
print(f"Total slides: {len(prs.slides)}")
for idx, slide in enumerate(prs.slides):
    print(f"\n=================== SLIDE {idx+1} ===================")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if p.text.strip():
                    print(f"- {p.text.strip()}")
        elif shape.has_table:
            for row in shape.table.rows:
                print(" | ".join(cell.text.strip() for cell in row.cells))
        else:
            print(f"[SHAPE]: {shape.name}")
