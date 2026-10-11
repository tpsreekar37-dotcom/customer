import pptx
import sys
sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation('Customer_Churn_Review_Presentation.pptx')
print(f"Total slides: {len(prs.slides)}\n")
for idx, slide in enumerate(prs.slides):
    titles = []
    for s in slide.shapes:
        if s.has_text_frame and s.text_frame.text:
            titles.append(s.text_frame.text.split('\n')[0].strip())
    print(f"Slide {idx+1:2d}: {titles[0] if titles else '[Visual Slide with Diagrams/Tables]'}")
